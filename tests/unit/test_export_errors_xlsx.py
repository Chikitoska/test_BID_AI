"""Unit/smoke: сбор строк XLSX без живого Influx."""

from __future__ import annotations

from datetime import datetime, timezone
from types import SimpleNamespace

from openpyxl import load_workbook
from io import BytesIO

from monitor.export_errors_xlsx import (
    ErrorRow,
    build_failure_flux,
    build_xlsx_bytes,
    normalize_run_type,
    parse_grafana_time,
    resolve_time_range,
    rows_from_records,
)


def test_normalize_run_type_maps_full_to_daily() -> None:
    assert normalize_run_type("health") == "health"
    assert normalize_run_type("full") == "daily"
    assert normalize_run_type("daily") == "daily"
    assert normalize_run_type("lk_pytest") == "lk_pytest"
    assert normalize_run_type("probe") == "probe"
    assert normalize_run_type("") == "unknown"


def test_parse_grafana_time_ms_and_iso() -> None:
    # 2024-06-15 12:00:00 UTC in ms
    ms = "1718452800000"
    dt = parse_grafana_time(ms)
    assert dt == datetime(2024, 6, 15, 12, 0, 0, tzinfo=timezone.utc)

    iso = parse_grafana_time("2024-06-15T12:00:00Z")
    assert iso == datetime(2024, 6, 15, 12, 0, 0, tzinfo=timezone.utc)

    sec = parse_grafana_time("1718452800")
    assert sec == datetime(2024, 6, 15, 12, 0, 0, tzinfo=timezone.utc)


def test_resolve_time_range_defaults_and_swap() -> None:
    start, stop = resolve_time_range(None, None, default_hours=1)
    assert (stop - start).total_seconds() == 3600

    a = "1718452800000"  # earlier
    b = "1718539200000"  # +24h
    start, stop = resolve_time_range(b, a)  # swapped inputs
    assert start < stop


def test_rows_from_records_mock_influx() -> None:
    t1 = datetime(2024, 6, 15, 10, 0, 0, tzinfo=timezone.utc)
    t2 = datetime(2024, 6, 15, 11, 0, 0, tzinfo=timezone.utc)

    class Rec:
        def __init__(self, ts, run_type, error):
            self._ts = ts
            self.values = {"run_type": run_type}
            self._error = error

        def get_time(self):
            return self._ts

        def get_value(self):
            return self._error

    rows = rows_from_records(
        [
            Rec(t1, "full", "HTTP 502"),
            Rec(t2, "health", "timeout"),
            Rec(t1, "lk_pytest", "assert failed"),
        ]
    )
    assert len(rows) == 3
    assert rows[0].date == t2  # newest first
    assert rows[0].run == "health"
    assert {r.run for r in rows} == {"health", "daily", "lk_pytest"}
    assert any(r.error == "HTTP 502" and r.run == "daily" for r in rows)


def test_build_xlsx_columns() -> None:
    rows = [
        ErrorRow(
            date=datetime(2024, 6, 15, 12, 0, 0, tzinfo=timezone.utc),
            error="HTTP 500",
            run="health",
        ),
        ErrorRow(
            date=datetime(2024, 6, 15, 11, 0, 0, tzinfo=timezone.utc),
            error="pytest fail",
            run="daily",
        ),
    ]
    raw = build_xlsx_bytes(rows)
    assert raw[:2] == b"PK"  # zip/xlsx magic
    wb = load_workbook(BytesIO(raw))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == ["дата", "ошибка", "прогон"]
    assert ws[2][1].value == "HTTP 500"
    assert ws[2][2].value == "health"
    assert ws[3][2].value == "daily"


def test_build_failure_flux_contains_measurement() -> None:
    start = datetime(2024, 1, 1, tzinfo=timezone.utc)
    stop = datetime(2024, 1, 2, tzinfo=timezone.utc)
    flux = build_failure_flux(start=start, stop=stop, bucket="bid_monitor")
    assert 'r._measurement == "bid_failure"' in flux
    assert "bid_monitor" in flux
    assert start.isoformat() in flux


def test_fetch_failure_rows_with_mock_query_api() -> None:
    from monitor.export_errors_xlsx import fetch_failure_rows

    t = datetime(2024, 6, 15, 12, 0, 0, tzinfo=timezone.utc)
    rec = SimpleNamespace(
        values={"run_type": "lk_pytest"},
        get_time=lambda: t,
        get_value=lambda: "UI timeout",
    )

    class Table:
        records = [rec]

    class QueryApi:
        def query(self, flux, org=None):
            assert "bid_failure" in flux
            return [Table()]

    rows = fetch_failure_rows(
        start=t,
        stop=t,
        query_api=QueryApi(),
    )
    assert len(rows) == 1
    assert rows[0].run == "lk_pytest"
    assert rows[0].error == "UI timeout"
